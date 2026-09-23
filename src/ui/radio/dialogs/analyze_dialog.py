"""
e-CALLISTO FITS Analyzer
Version 3.1.0
Sahan S Liyanage (sahanslst@gmail.com)
Astronomical and Space Science Unit, University of Colombo, Sri Lanka.
"""

from __future__ import annotations

import os
import re
import sys

import numpy as np
from PySide6.QtCore import Qt, QTime, Signal
from PySide6.QtGui import QFont
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QDialog,
    QDoubleSpinBox,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QStatusBar,
    QTableWidget,
    QTableWidgetItem,
    QTimeEdit,
    QVBoxLayout,
    QWidget,
)

from src.backend.common.figure_export import FIGURE_EXPORT_FILTERS, save_figure
from src.backend.radio.density_models import (
    DEFAULT_DENSITY_MODEL,
    DENSITY_MODEL_ORDER,
    DENSITY_MODELS,
    density_model_label,
    normalize_density_model,
    shock_parameters,
)
from src.backend.radio.radio_figures import draw_fit_graph, fit_graph_figure, origin_graph_style, power_law_label
from src.backend.session.analysis_session import DEFAULT_T0_MODE, normalize_t0_mode
from src.ui.common.gui_shared import MplCanvas, fit_window_to_screen, pick_export_path
from src.ui.common.mpl_style import style_axes

#: t0 choices offered by the Analyzer: (session value, label).
T0_CHOICES = (
    ("file_start", "File start"),
    ("burst_onset", "Burst onset"),
    ("custom", "Custom time"),
)
#: Columns of the density-model comparison table: (summary key, header, tooltip, digits).
_COMPARISON_COLUMNS = (
    ("initial_shock_speed_km_s", "v₀ (km/s)", "Initial shock speed", 1),
    ("initial_shock_height_rs", "h₀ (R☉)", "Initial shock height", 3),
    ("avg_shock_speed_km_s", "v̄ (km/s)", "Average shock speed", 1),
    ("avg_shock_height_rs", "h̄ (R☉)", "Average shock height", 3),
)


class AnalyzeDialog(QDialog):
    sessionChanged = Signal(dict)

    def __init__(
        self,
        time_channels,
        freqs,
        filename,
        fundamental=True,
        harmonic=False,
        parent=None,
        session=None,
        time_seconds=None,
    ):
        super().__init__(parent)
        self.fundamental = fundamental
        self.harmonic = harmonic

        self.setWindowTitle("Analyzer")
        fit_window_to_screen(self, 1100, 700)

        self.time_channels = np.asarray(time_channels, dtype=float).reshape(-1)
        self.time_seconds = self._resolve_time_seconds(self.time_channels, time_seconds)
        self.time = self._build_analysis_time_axis(self.time_seconds)
        self.freq = np.asarray(freqs, dtype=float).reshape(-1)
        self.filename = filename.split(".")[0]
        self.current_plot_title = f"{self.filename}_Best_Fit"
        #: What the canvas shows, as fit_graph_figure() arguments: "Save Graph" redraws it Origin-style.
        self._graph = None
        self._fit_params = None
        self._shock_summary = {}
        self._type_ii_state = None
        self._suppress_emit = False
        self._fit_mask = np.isfinite(self.time) & np.isfinite(self.freq)
        #: UT of the first data sample (seconds of day), for a custom t0 in UT.
        self._ut_start_sec = None
        #: The t0 (seconds from file start) the current fit was made with.
        self._fit_t0_s = 0.0
        self._model_comparison = {}

        # Canvas
        self.canvas = MplCanvas(self, width=8, height=5)
        self.canvas.figure.clf()
        self.canvas.ax = self.canvas.figure.add_subplot(111)
        style_axes(self.canvas.ax)

        # Buttons
        self.max_button = QPushButton("Maximum Intensities")
        self.fit_button = QPushButton("Best Fit")
        self.save_plot_button = QPushButton("Save Graph")
        self.save_data_button = QPushButton("Save Data")
        self.existing_excel_checkbox = QCheckBox("Existing Excel File")

        self.extra_plot_label = QLabel("Extra Plots:")
        self.extra_plot_combo = QComboBox()
        self.extra_plot_combo.addItems([
            "Shock Speed vs Shock Height",
            "Shock Speed vs Frequency",
            "Shock Height vs Frequency"
        ])
        self.extra_plot_button = QPushButton("Plot")

        self.max_button.clicked.connect(self.plot_max)
        self.fit_button.clicked.connect(self.plot_fit)
        self.save_plot_button.clicked.connect(self.save_graph)
        self.save_data_button.clicked.connect(self.save_data)
        self.extra_plot_button.clicked.connect(self.plot_extra)

        # Plot control layout
        plot_button_layout = QHBoxLayout()
        plot_button_layout.addWidget(self.max_button)
        plot_button_layout.addWidget(self.fit_button)

        # --- Power-law time origin t0: f(t) = a (t - t0)^-b ---
        self.t0_label = QLabel("t₀:")
        self.t0_combo = QComboBox()
        for value, label in T0_CHOICES:
            self.t0_combo.addItem(label, value)
        self.t0_combo.setToolTip(
            "Time origin of the power-law fit f = a·(t − t₀)^−b.\n"
            "File start: t measured from the start of the loaded data (the previous behaviour).\n"
            "Burst onset: one sample before the first fitted point.\n"
            "Custom time: a time you enter."
        )
        self.t0_time_edit = QTimeEdit()
        self.t0_time_edit.setDisplayFormat("HH:mm:ss.zzz")
        self.t0_time_edit.setToolTip("Custom t₀ in UT")
        self.t0_seconds_spin = QDoubleSpinBox()
        self.t0_seconds_spin.setDecimals(2)
        self.t0_seconds_spin.setRange(-86400.0, 86400.0)
        self.t0_seconds_spin.setSuffix(" s")
        self.t0_seconds_spin.setToolTip("Custom t₀ in seconds from the start of the loaded data")
        self.t0_combo.currentIndexChanged.connect(self._on_t0_mode_changed)
        self.t0_time_edit.editingFinished.connect(self._on_t0_value_edited)
        self.t0_seconds_spin.editingFinished.connect(self._on_t0_value_edited)
        plot_button_layout.addSpacing(12)
        plot_button_layout.addWidget(self.t0_label)
        plot_button_layout.addWidget(self.t0_combo)
        plot_button_layout.addWidget(self.t0_time_edit)
        plot_button_layout.addWidget(self.t0_seconds_spin)
        plot_button_layout.addStretch(1)

        left_layout = QVBoxLayout()
        left_layout.addLayout(plot_button_layout)
        left_layout.addWidget(self.canvas)

        # === Info Panel ===

        # --- Newkirk fold selection (n-fold) ---
        self.fold_label = QLabel("Fold-number:")
        self.fold_combo = QComboBox()
        self.fold_combo.addItems(["1", "2", "3", "4"])
        self.fold_combo.setCurrentIndex(0)

        self.fold_calc_button = QPushButton("Calculate")
        self.fold_calc_button.setEnabled(False)  # enable only after Best Fit
        self.fold_calc_button.clicked.connect(self.recalculate_shock_parameters)

        # Put fold controls into a widget so it can live inside self.labels
        self.fold_row_widget = QWidget()
        fold_row_layout = QHBoxLayout(self.fold_row_widget)
        fold_row_layout.setContentsMargins(0, 0, 0, 0)
        fold_row_layout.addWidget(self.fold_label)
        fold_row_layout.addWidget(self.fold_combo)
        fold_row_layout.addWidget(self.fold_calc_button)

        # Reserve enough room for the selected fold value across Qt styles/themes.
        self.fold_combo.setMinimumContentsLength(2)
        self.fold_combo.setSizeAdjustPolicy(QComboBox.SizeAdjustPolicy.AdjustToMinimumContentsLengthWithIcon)
        self.fold_combo.setMinimumWidth(max(70, self.fold_combo.sizeHint().width()))

        # --- Coronal density model ---
        self.density_model_label = QLabel("Density model:")
        self.density_model_combo = QComboBox()
        for key in DENSITY_MODEL_ORDER:
            model = DENSITY_MODELS[key]
            self.density_model_combo.addItem(model.label, key)
            self.density_model_combo.setItemData(
                self.density_model_combo.count() - 1, model.reference, Qt.ItemDataRole.ToolTipRole
            )
        self.density_model_combo.setToolTip("Density model used to turn frequency and drift into shock height and speed.")
        self.density_model_combo.currentIndexChanged.connect(self._on_density_model_changed)
        self.model_row_widget = QWidget()
        model_row_layout = QHBoxLayout(self.model_row_widget)
        model_row_layout.setContentsMargins(0, 0, 0, 0)
        model_row_layout.addWidget(self.density_model_label)
        model_row_layout.addWidget(self.density_model_combo, 1)

        # --- Side-by-side comparison of every density model ---
        self.comparison_header = QLabel("<b>Density model comparison:</b>")
        self.comparison_table = QTableWidget(len(DENSITY_MODEL_ORDER), len(_COMPARISON_COLUMNS))
        self.comparison_table.setVerticalHeaderLabels([DENSITY_MODELS[key].label for key in DENSITY_MODEL_ORDER])
        for column, (_key, header, tooltip, _digits) in enumerate(_COMPARISON_COLUMNS):
            item = QTableWidgetItem(header)
            item.setToolTip(tooltip)
            self.comparison_table.setHorizontalHeaderItem(column, item)
        self.comparison_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.comparison_table.setSelectionMode(QAbstractItemView.SelectionMode.NoSelection)
        self.comparison_table.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        self.comparison_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.comparison_table.verticalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.comparison_table.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.comparison_table.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self._fit_comparison_table_height()

        self.equation_label = QLabel("Best Fit Equation:")
        self.equation_display = QLabel("")
        self.equation_display.setTextFormat(Qt.RichText)
        self.equation_display.setStyleSheet("font-size: 16px; padding: 4px;")

        self.stats_header = QLabel("<b>Fit Metrics:</b>")
        self.r2_display = QLabel("R² = ")
        self.rmse_display = QLabel("RMSE = ")

        self.shock_header = QLabel("<b>Shock Parameters:</b>")
        self.avg_freq_display = QLabel("")
        self.drift_display = QLabel("")
        self.start_freq_display = QLabel("")
        self.initial_shock_speed_display = QLabel("")
        self.initial_shock_height_display = QLabel("")
        self.avg_shock_speed_display = QLabel("")
        self.avg_shock_height_display = QLabel("")

        self.labels = [
            self.model_row_widget,
            self.fold_row_widget,
            self.equation_label, self.equation_display,
            self.stats_header, self.r2_display, self.rmse_display,
            self.shock_header,
            self.avg_freq_display, self.drift_display, self.start_freq_display,
            self.initial_shock_speed_display, self.initial_shock_height_display,
            self.avg_shock_speed_display, self.avg_shock_height_display,
            self.comparison_header, self.comparison_table,
            self.save_plot_button, self.save_data_button, self.existing_excel_checkbox,
            self.extra_plot_label, self.extra_plot_combo, self.extra_plot_button
        ]

        right_inner = QVBoxLayout()
        for widget in self.labels:
            right_inner.addWidget(widget)
        right_inner.addStretch()

        right_inner.addStretch()

        right_widget = QWidget()
        right_widget.setLayout(right_inner)
        right_scroll = QScrollArea()
        right_scroll.setWidget(right_widget)
        right_scroll.setWidgetResizable(True)
        right_scroll.setMinimumWidth(300)

        self.status = QStatusBar()
        self.status.showMessage("Ready")

        # Main layout
        main_layout = QHBoxLayout()
        main_layout.addLayout(left_layout, stretch=3)
        main_layout.addWidget(right_scroll, stretch=1)
        main_with_status = QVBoxLayout()
        main_with_status.addLayout(main_layout)
        main_with_status.addWidget(self.status)
        self.setLayout(main_with_status)
        self._sync_t0_widgets()
        self._refresh_comparison_table()

        if isinstance(session, dict):
            try:
                self.restore_session(session, emit_change=False)
            except Exception:
                pass

    @staticmethod
    def _resolve_time_seconds(time_channels, time_seconds=None):
        channels = np.asarray(time_channels, dtype=float).reshape(-1)
        if time_seconds is not None:
            arr = np.asarray(time_seconds, dtype=float).reshape(-1)
            if arr.shape == channels.shape:
                return arr
        return channels * 0.25

    def _build_analysis_time_axis(self, time_seconds):
        arr = np.asarray(time_seconds, dtype=float).reshape(-1)
        if arr.size == 0:
            return arr
        finite = np.isfinite(arr)
        if not np.any(finite):
            return np.arange(arr.size, dtype=float)
        return np.where(finite, arr, np.nan)

    @staticmethod
    def _initial_power_law_guess(time_values, freq_values):
        x = np.asarray(time_values, dtype=float).reshape(-1)
        y = np.asarray(freq_values, dtype=float).reshape(-1)
        mask = np.isfinite(x) & np.isfinite(y) & (x > 0) & (y > 0)
        if np.count_nonzero(mask) >= 2:
            lx = np.log(x[mask])
            ly = np.log(y[mask])
            try:
                slope, intercept = np.polyfit(lx, ly, 1)
                a0 = float(np.exp(intercept))
                b0 = float(max(1e-9, -slope))
                if np.isfinite(a0) and a0 > 0 and np.isfinite(b0):
                    return a0, b0
            except Exception:
                pass

        positive_freqs = y[np.isfinite(y) & (y > 0)]
        if positive_freqs.size:
            return float(np.nanmax(positive_freqs)), 0.5
        return 1.0, 0.5

    @staticmethod
    def _power_law_fit_mask(time_values, freq_values):
        x = np.asarray(time_values, dtype=float).reshape(-1)
        y = np.asarray(freq_values, dtype=float).reshape(-1)
        return np.isfinite(x) & np.isfinite(y) & (x > 0) & (y > 0)

    # ---- power-law time origin t0 -------------------------------------------
    def set_ut_start_sec(self, ut_start_sec) -> None:
        """UT (seconds of day) of the first data sample, so a custom t₀ can be entered in UT."""
        try:
            value = float(ut_start_sec)
        except (TypeError, ValueError):
            value = None
        if value is not None and not np.isfinite(value):
            value = None
        custom_t0 = self._custom_t0_seconds()
        self._ut_start_sec = value
        self._set_custom_t0_seconds(custom_t0)
        self._sync_t0_widgets()

    def _t0_mode(self) -> str:
        return normalize_t0_mode(self.t0_combo.currentData())

    def _set_t0_mode(self, mode) -> None:
        index = self.t0_combo.findData(normalize_t0_mode(mode))
        blocked = self.t0_combo.blockSignals(True)
        try:
            self.t0_combo.setCurrentIndex(max(0, index))
        finally:
            self.t0_combo.blockSignals(blocked)
        self._sync_t0_widgets()

    def _custom_t0_seconds(self) -> float:
        """The custom t₀ in seconds from the start of the loaded data."""
        if self._ut_start_sec is not None:
            seconds_of_day = self.t0_time_edit.time().msecsSinceStartOfDay() / 1000.0
            offset = seconds_of_day - float(self._ut_start_sec)
            # A burst that runs past midnight UT.
            if offset < -43200.0:
                offset += 86400.0
            elif offset > 43200.0:
                offset -= 86400.0
            return float(offset)
        return float(self.t0_seconds_spin.value())

    def _set_custom_t0_seconds(self, seconds) -> None:
        try:
            value = float(seconds)
        except (TypeError, ValueError):
            value = 0.0
        if not np.isfinite(value):
            value = 0.0
        for widget in (self.t0_seconds_spin, self.t0_time_edit):
            widget.blockSignals(True)
        try:
            self.t0_seconds_spin.setValue(value)
            if self._ut_start_sec is not None:
                msecs = int(round(((float(self._ut_start_sec) + value) % 86400.0) * 1000.0))
                self.t0_time_edit.setTime(QTime(0, 0).addMSecs(msecs))
        finally:
            for widget in (self.t0_seconds_spin, self.t0_time_edit):
                widget.blockSignals(False)

    def _burst_onset_t0(self) -> float:
        """One sample before the first point, so the power law stays finite there."""
        times = np.asarray(self.time, dtype=float).reshape(-1)
        freqs = np.asarray(self.freq, dtype=float).reshape(-1)
        if times.shape != freqs.shape:
            return 0.0
        valid = np.isfinite(times) & np.isfinite(freqs) & (freqs > 0.0)
        if not np.any(valid):
            return 0.0
        ordered = np.unique(times[valid])
        steps = np.diff(ordered)
        steps = steps[steps > 0.0]
        step = float(np.median(steps)) if steps.size else 0.25
        return float(ordered[0] - step)

    def _resolved_t0(self) -> float:
        mode = self._t0_mode()
        if mode == "burst_onset":
            return self._burst_onset_t0()
        if mode == "custom":
            return self._custom_t0_seconds()
        return 0.0

    def _sync_t0_widgets(self) -> None:
        custom = self._t0_mode() == "custom"
        self.t0_time_edit.setVisible(custom and self._ut_start_sec is not None)
        self.t0_seconds_spin.setVisible(custom and self._ut_start_sec is None)

    def _on_t0_mode_changed(self, _index=0):
        if self._t0_mode() == "custom" and not getattr(self, "_custom_t0_initialized", False):
            # Start the custom value somewhere sensible: the burst onset.
            self._set_custom_t0_seconds(self._burst_onset_t0())
            self._custom_t0_initialized = True
        self._sync_t0_widgets()
        self._refit_after_t0_change()

    def _on_t0_value_edited(self):
        if self._t0_mode() == "custom":
            self._refit_after_t0_change()

    def _refit_after_t0_change(self):
        """A new t₀ changes the fit itself, so an existing Best Fit is redone."""
        if self._suppress_emit or self._fit_params is None:
            self._emit_session_changed()
            return
        if abs(self._resolved_t0() - float(self._fit_t0_s)) < 1e-9:
            return
        self.plot_fit()

    # ---- density model -------------------------------------------------------
    def _selected_density_model(self) -> str:
        return normalize_density_model(self.density_model_combo.currentData())

    def _set_density_model(self, model) -> None:
        index = self.density_model_combo.findData(normalize_density_model(model))
        blocked = self.density_model_combo.blockSignals(True)
        try:
            self.density_model_combo.setCurrentIndex(max(0, index))
        finally:
            self.density_model_combo.blockSignals(blocked)

    def _on_density_model_changed(self, _index=0):
        if self._suppress_emit:
            return
        if hasattr(self, "_drift_vals"):
            self._update_shock_parameters(self._selected_fold())
            label = density_model_label(self._selected_density_model())
            self.status.showMessage(f"Updated using the {label} density model.", 3000)
        else:
            self._refresh_comparison_table()
        self._emit_session_changed()

    def _fit_comparison_table_height(self) -> None:
        table = self.comparison_table
        header = table.horizontalHeader().sizeHint().height()
        rows = sum(table.rowHeight(row) for row in range(table.rowCount()))
        table.setFixedHeight(header + rows + 2 * table.frameWidth() + 2)

    def _refresh_comparison_table(self) -> None:
        selected = self._selected_density_model()
        have_results = bool(self._model_comparison)
        for row, key in enumerate(DENSITY_MODEL_ORDER):
            result = self._model_comparison.get(key) or {}
            for column, (field, _header, _tip, digits) in enumerate(_COMPARISON_COLUMNS):
                value = result.get(field)
                text, tooltip = "—", ""
                if value is not None and np.isfinite(value):
                    text = f"{float(value):.{digits}f}"
                elif have_results:
                    tooltip = "This frequency range lies outside the model (below the photosphere or beyond 1 AU)."
                item = QTableWidgetItem(text)
                item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                if tooltip:
                    item.setToolTip(tooltip)
                if key == selected:
                    font = QFont(item.font())
                    font.setBold(True)
                    item.setFont(font)
                self.comparison_table.setItem(row, column, item)
            header_item = self.comparison_table.verticalHeaderItem(row) or QTableWidgetItem(DENSITY_MODELS[key].label)
            header_font = QFont(header_item.font())
            header_font.setBold(key == selected)
            header_item.setFont(header_font)
            self.comparison_table.setVerticalHeaderItem(row, header_item)
        self.comparison_header.setText(f"<b>Density model comparison ({self._selected_fold()}-fold):</b>")
        self._fit_comparison_table_height()

    def _emit_session_changed(self):
        if self._suppress_emit:
            return
        try:
            self.sessionChanged.emit(self.session_state())
        except Exception:
            pass

    def _set_summary_labels_from_dict(self, summary: dict):
        if not isinstance(summary, dict):
            return
        fold = int(summary.get("fold", self._selected_fold()) or self._selected_fold())
        model = density_model_label(summary.get("density_model", self._selected_density_model()))
        self.shock_header.setText(f"<b>Shock Parameters ({model} {fold}-fold):</b>")

        def _f(v, digits=2):
            if v is None:
                return ""
            try:
                value = float(v)
            except Exception:
                return ""
            # A frequency outside the chosen density model's range has no value.
            return f"{value:.{digits}f}" if np.isfinite(value) else "—"

        self.avg_freq_display.setText(
            f"Average Frequency: <b>{_f(summary.get('avg_freq_mhz'), 2)} ± {_f(summary.get('avg_freq_err_mhz'), 2)}</b> MHz"
        )
        self.drift_display.setText(
            f"Average Drift Rate: <b>{_f(summary.get('avg_drift_mhz_s'), 4)} ± {_f(summary.get('avg_drift_err_mhz_s'), 4)}</b> MHz/s"
        )
        self.start_freq_display.setText(
            f"Starting Frequency: <b>{_f(summary.get('start_freq_mhz'), 2)} ± {_f(summary.get('start_freq_err_mhz'), 2)}</b> MHz"
        )
        self.initial_shock_speed_display.setText(
            f"Initial Shock Speed: <b>{_f(summary.get('initial_shock_speed_km_s'), 2)} ± {_f(summary.get('initial_shock_speed_err_km_s'), 2)}</b> km/s"
        )
        self.initial_shock_height_display.setText(
            f"Initial Shock Height: <b>{_f(summary.get('initial_shock_height_rs'), 3)} ± {_f(summary.get('initial_shock_height_err_rs'), 3)}</b> Rₛ"
        )
        self.avg_shock_speed_display.setText(
            f"Average Shock Speed: <b>{_f(summary.get('avg_shock_speed_km_s'), 2)} ± {_f(summary.get('avg_shock_speed_err_km_s'), 2)}</b> km/s"
        )
        self.avg_shock_height_display.setText(
            f"Average Shock Height: <b>{_f(summary.get('avg_shock_height_rs'), 3)} ± {_f(summary.get('avg_shock_height_err_rs'), 3)}</b> Rₛ"
        )

    def plot_max(self):
        self.canvas.ax.clear()
        self.canvas.ax.scatter(self.time, self.freq, s=10, color='blue')
        self.canvas.ax.set_title(f"{self.filename}_Maximum_Intensity")
        self.canvas.ax.set_xlabel("Time (s)")
        self.canvas.ax.set_ylabel("Frequency (MHz)")
        self.canvas.ax.grid(True)
        self.canvas.draw()
        self.current_plot_title = f"{self.filename}_Maximum_Intensity"
        self._graph = {
            "x": self.time,
            "y": self.freq,
            "data_label": "",
            "title": self.current_plot_title,
            "x_label": "Time (s)",
            "y_label": "Frequency (MHz)",
        }
        self.equation_display.setText("")
        self.fold_calc_button.setEnabled(False)
        self.status.showMessage("Max intensities plotted successfully!", 3000)

    def plot_fit(self, _checked=False, params=None, std_errs=None):
        def model_func(x, a, b): return a * x ** (-b)

        def drift_rate(x, a_, b_): return -a_ * b_ * x ** (-b_ - 1)

        plot_time = np.asarray(self.time, dtype=float).reshape(-1)
        plot_freq = np.asarray(self.freq, dtype=float).reshape(-1)
        # The power law is fitted in time since t0; the graph keeps the data's own time axis.
        t0 = self._resolved_t0()
        fit_mask = self._power_law_fit_mask(plot_time - t0, plot_freq)
        if np.count_nonzero(fit_mask) < 2:
            where = "time > 0 s" if t0 == 0.0 else f"time after t₀ = {t0:.2f} s"
            QMessageBox.warning(
                self,
                "Analyzer",
                f"Power-law fitting requires at least two points with {where} and frequency > 0 MHz.",
            )
            self.status.showMessage("Best fit failed: insufficient positive time samples.", 3000)
            return

        fit_time = plot_time[fit_mask] - t0
        fit_freq = plot_freq[fit_mask]
        self._fit_mask = fit_mask
        self._fit_t0_s = float(t0)

        if params is None:
            from scipy.optimize import curve_fit

            params, cov = curve_fit(
                model_func,
                fit_time,
                fit_freq,
                p0=self._initial_power_law_guess(fit_time, fit_freq),
                bounds=([1e-12, 1e-9], [np.inf, np.inf]),
                maxfev=10000,
            )
            a, b = params
            std_errs = np.sqrt(np.diag(cov))
        else:
            a, b = float(params[0]), abs(float(params[1]))
            if std_errs is None:
                std_errs = np.array([np.nan, np.nan], dtype=float)

        time_fit = np.linspace(fit_time.min(), fit_time.max(), 400)
        freq_fit = model_func(time_fit, a, b)

        # The best fit is the one Analyzer graph shown in the OriginPro style
        # on screen; "Save Graph" draws the same graph on its own page.
        self.current_plot_title = f"{self.filename}_Best_Fit"
        variable = "x" if t0 == 0.0 else f"(x − {t0:.2f})"
        self._graph = {
            "x": plot_time,
            "y": plot_freq,
            "fit_x": time_fit + t0,
            "fit_y": freq_fit,
            "data_label": "Original Data",
            "fit_label": power_law_label(a, b, prefix="Best Fit", variable=variable),
            "title": self.current_plot_title,
            "x_label": "Time (s)",
            "y_label": "Frequency (MHz)",
        }
        self.canvas.ax.clear()
        self.canvas.figure.clf()
        with origin_graph_style() as text:
            self.canvas.figure.set_facecolor("white")
            self.canvas.ax = self.canvas.figure.add_subplot(111)
            draw_fit_graph(self.canvas.ax, **self._graph, text=text)
            self.canvas.draw()

        from sklearn.metrics import mean_squared_error, r2_score

        predicted = model_func(fit_time, a, b)
        r2 = r2_score(fit_freq, predicted)
        rmse = np.sqrt(mean_squared_error(fit_freq, predicted))

        if t0 == 0.0:
            self.equation_display.setText(f"<b>f(x) = {a:.2f} · x<sup>-{b:.2f}</sup></b>")
        else:
            self.equation_display.setText(f"<b>f(x) = {a:.2f} · (x − {t0:.2f})<sup>-{b:.2f}</sup></b>")
        self.r2_display.setText(f"R² = {r2:.4f}")
        self.rmse_display.setText(f"RMSE = {rmse:.4f}")

        # Cache fit parameters for session persistence
        try:
            self._fit_params = {
                "a": float(a),
                "b": float(b),
                "std_errs": [float(std_errs[0]), float(std_errs[1])],
                "r2": float(r2),
                "rmse": float(rmse),
            }
        except Exception:
            self._fit_params = None

        drift_vals = drift_rate(fit_time, a, b)
        shock_calc_time = time_fit
        shock_calc_freq = freq_fit
        shock_calc_drift_vals = drift_rate(shock_calc_time, a, b)
        residuals = fit_freq - predicted
        freq_err = np.std(residuals)
        b_err_denom = max(abs(float(b)), 1e-9)
        drift_errs = np.abs(drift_vals) * np.sqrt((std_errs[0] / a) ** 2 + (std_errs[1] / b_err_denom) ** 2)
        shock_calc_drift_errs = np.abs(shock_calc_drift_vals) * np.sqrt(
            (std_errs[0] / a) ** 2 + (std_errs[1] / b_err_denom) ** 2
        )

        # Cache results so we can recompute shock params for different folds
        self._fit_time = fit_time
        self._fit_freq = fit_freq
        self._drift_vals = drift_vals
        self._drift_errs = drift_errs
        self._shock_calc_time = shock_calc_time
        self._shock_calc_freq = shock_calc_freq
        self._shock_calc_drift_vals = shock_calc_drift_vals
        self._shock_calc_drift_errs = shock_calc_drift_errs
        self.freq_err = freq_err

        # Enable fold recalculation now that Best Fit exists
        self.fold_calc_button.setEnabled(True)

        # Compute and display shock parameters using selected fold-number
        self._update_shock_parameters(self._selected_fold())

        self.status.showMessage("Best fit plotted successfully!", 3000)
        self._emit_session_changed()

    def session_state(self) -> dict:
        fold = int(self._selected_fold())
        fit = dict(getattr(self, "_fit_params", None) or {})
        shock = dict(getattr(self, "_shock_summary", {}) or {})
        if shock:
            shock["fold"] = int(shock.get("fold", fold) or fold)
            shock["fundamental"] = bool(shock.get("fundamental", self.fundamental))
            shock["harmonic"] = bool(shock.get("harmonic", self.harmonic))

        return {
            "source": {"filename": str(self.filename or "")},
            "max_intensity": {
                "time_channels": np.asarray(self.time_channels, dtype=float),
                "time_seconds": np.asarray(self.time_seconds, dtype=float),
                "freqs": np.asarray(self.freq, dtype=float),
                "fundamental": bool(self.fundamental),
                "harmonic": bool(self.harmonic),
            },
            "analyzer": {
                "fit_params": fit,
                "fold": fold,
                "shock_summary": shock,
                "density_model": self._selected_density_model(),
                "t0_mode": self._t0_mode(),
                "t0_s": float(self._fit_t0_s if fit else self._resolved_t0()),
            },
            "type_ii": dict(self._type_ii_state or {}),
            "ui": {
                "restore_max_window": True,
                "restore_analyzer_window": True,
                "restore_type_ii_window": bool(self._type_ii_state),
            },
        }

    def restore_session(self, state: dict, *, emit_change: bool = True):
        self._suppress_emit = True
        try:
            max_block = dict(state.get("max_intensity") or state) if isinstance(state, dict) else {}
            analyzer = dict(state.get("analyzer") or state) if isinstance(state, dict) else {}
            type_ii = dict(state.get("type_ii") or {}) if isinstance(state, dict) else {}
            fit = analyzer.get("fit_params", None)
            shock = analyzer.get("shock_summary", None)

            if max_block.get("time_channels") is not None and max_block.get("freqs") is not None:
                try:
                    t_arr = np.asarray(max_block.get("time_channels"), dtype=float).reshape(-1)
                    ts_arr = max_block.get("time_seconds", None)
                    f_arr = np.asarray(max_block.get("freqs"), dtype=float).reshape(-1)
                    if len(t_arr) == len(f_arr) and len(t_arr) > 0:
                        self.time_channels = t_arr
                        self.time_seconds = self._resolve_time_seconds(self.time_channels, ts_arr)
                        self.time = self._build_analysis_time_axis(self.time_seconds)
                        self.freq = f_arr
                        self._fit_mask = np.isfinite(self.time) & np.isfinite(self.freq)
                except Exception:
                    pass

            self.fundamental = bool(max_block.get("fundamental", self.fundamental))
            self.harmonic = bool(max_block.get("harmonic", self.harmonic))

            try:
                fold = int(analyzer.get("fold", 1))
            except Exception:
                fold = 1
            fold = max(1, min(4, fold))
            try:
                self.fold_combo.setCurrentIndex(fold - 1)
            except Exception:
                pass

            self._set_density_model(analyzer.get("density_model", DEFAULT_DENSITY_MODEL))
            t0_mode = normalize_t0_mode(analyzer.get("t0_mode", DEFAULT_T0_MODE))
            if t0_mode == "custom":
                self._set_custom_t0_seconds(analyzer.get("t0_s", 0.0))
                self._custom_t0_initialized = True
            self._set_t0_mode(t0_mode)

            if not isinstance(fit, dict):
                fit = None

            if fit is not None and ("a" in fit and "b" in fit):
                try:
                    a = float(fit["a"])
                    b = float(fit["b"])
                except Exception:
                    a = None
                    b = None

                if a is not None and b is not None:
                    std_errs = fit.get("std_errs", None)
                    std_errs_arr = None
                    if isinstance(std_errs, (list, tuple)) and len(std_errs) >= 2:
                        try:
                            std_errs_arr = np.array([float(std_errs[0]), float(std_errs[1])], dtype=float)
                        except Exception:
                            std_errs_arr = None
                    self.plot_fit(params=(a, b), std_errs=std_errs_arr)
            elif isinstance(shock, dict):
                self._shock_summary = dict(shock)
                self._set_summary_labels_from_dict(self._shock_summary)

            self._type_ii_state = dict(type_ii) if isinstance(type_ii, dict) and type_ii else None
        finally:
            self._suppress_emit = False
        if emit_change:
            self._emit_session_changed()

    def _selected_fold(self):
        try:
            n = int(self.fold_combo.currentText())
        except Exception:
            n = 1
        return max(1, min(4, n))

    def recalculate_shock_parameters(self):
        if not hasattr(self, "_drift_vals") or not hasattr(self, "_drift_errs"):
            QMessageBox.information(self, "Analyzer", "Please click 'Best Fit' first.")
            return

        n = self._selected_fold()
        self._update_shock_parameters(n)
        label = density_model_label(self._selected_density_model())
        self.status.showMessage(f"Updated using {label} {n}-fold model.", 3000)
        self._emit_session_changed()

    def _update_shock_parameters(self, n):
        observed_freq_values = np.asarray(getattr(self, "_fit_freq", self.freq), dtype=float).reshape(-1)
        observed_drift_vals = np.asarray(self._drift_vals, dtype=float).reshape(-1)
        max_intensity_freq_values = np.asarray(self.freq, dtype=float).reshape(-1)
        max_intensity_freq_values = max_intensity_freq_values[
            np.isfinite(max_intensity_freq_values) & (max_intensity_freq_values > 0.0)
        ]
        freq_values = np.asarray(getattr(self, "_shock_calc_freq", observed_freq_values), dtype=float).reshape(-1)
        drift_vals = np.asarray(getattr(self, "_shock_calc_drift_vals", observed_drift_vals), dtype=float).reshape(-1)
        drift_errs = np.asarray(getattr(self, "_shock_calc_drift_errs", self._drift_errs), dtype=float).reshape(-1)
        if freq_values.size == 0 or drift_vals.size == 0 or freq_values.size != drift_vals.size:
            return

        harmonic_number = 2.0 if self.harmonic else 1.0
        shock_freq_values = freq_values / harmonic_number
        shock_drift_vals = drift_vals / harmonic_number
        shock_drift_errs = drift_errs / harmonic_number
        shock_freq_err = float(self.freq_err) / harmonic_number

        # Every density model is evaluated, for the side-by-side table; the
        # selected one drives the labels, the session and the extra plots.
        comparison = {
            key: shock_parameters(
                shock_freq_values,
                shock_drift_vals,
                shock_drift_errs,
                freq_err_mhz=shock_freq_err,
                model=key,
                fold=n,
                start_percentile=90,
            )
            for key in DENSITY_MODEL_ORDER
        }
        model = self._selected_density_model()
        selected = comparison[model]
        shock_speed = selected["shock_speed_km_s"]
        R_p = selected["shock_height_rs"]
        avg_freq_values = max_intensity_freq_values / harmonic_number if max_intensity_freq_values.size else shock_freq_values

        percentile = 90
        start_freq = selected["start_freq_mhz"]
        start_shock_speed = selected["initial_shock_speed_km_s"]
        start_height = selected["initial_shock_height_rs"]
        shock_speed_err = selected["initial_shock_speed_err_km_s"]
        Rp_err = selected["initial_shock_height_err_rs"]

        # Average frequency is reported from the selected maximum-intensity points.
        avg_freq = np.mean(avg_freq_values)
        avg_freq_err = np.std(avg_freq_values) / np.sqrt(len(avg_freq_values))
        avg_drift = np.mean(shock_drift_vals)
        avg_drift_err = np.std(shock_drift_vals) / np.sqrt(len(shock_drift_vals))

        avg_speed = selected["avg_shock_speed_km_s"]
        avg_speed_err = selected["avg_shock_speed_err_km_s"]
        avg_height = selected["avg_shock_height_rs"]
        avg_height_err = selected["avg_shock_height_err_rs"]

        # Store arrays for extra plots
        self.shock_speed = shock_speed
        self.R_p = R_p
        self._shock_freq_values = shock_freq_values
        self._shock_drift_vals = shock_drift_vals
        self.start_freq = start_freq
        self.start_height = start_height

        self._shock_summary = {
            "avg_freq_mhz": float(avg_freq),
            "avg_freq_err_mhz": float(avg_freq_err),
            "avg_drift_mhz_s": float(avg_drift),
            "avg_drift_err_mhz_s": float(avg_drift_err),
            "start_freq_mhz": float(start_freq),
            "start_freq_err_mhz": float(shock_freq_err),
            "initial_shock_speed_km_s": float(start_shock_speed),
            "initial_shock_speed_err_km_s": float(shock_speed_err),
            "initial_shock_height_rs": float(start_height),
            "initial_shock_height_err_rs": float(Rp_err),
            "avg_shock_speed_km_s": float(avg_speed),
            "avg_shock_speed_err_km_s": float(avg_speed_err),
            "avg_shock_height_rs": float(avg_height),
            "avg_shock_height_err_rs": float(avg_height_err),
            "fold": int(n),
            "fundamental": bool(self.fundamental),
            "harmonic": bool(self.harmonic),
            "harmonic_number": int(harmonic_number),
            "observed_avg_freq_mhz": float(np.mean(max_intensity_freq_values)) if max_intensity_freq_values.size else float(np.mean(observed_freq_values)),
            "observed_avg_drift_mhz_s": float(np.mean(observed_drift_vals)),
            "observed_start_freq_mhz": float(np.percentile(observed_freq_values, percentile)),
            "density_model": model,
        }
        self._model_comparison = comparison

        self._set_summary_labels_from_dict(self._shock_summary)
        self._refresh_comparison_table()

    def save_graph(self):
        """Save the graph on show as an OriginPro-style figure; the window keeps its look."""
        if self._graph is None:
            QMessageBox.information(self, "Save Graph", "Plot the maximum intensities or the best fit first.")
            return

        plot_name = getattr(self, "current_plot_title", None) or f"{self.filename}_Plot"

        file_path, ext = pick_export_path(
            self,
            "Export Figure",
            plot_name,
            FIGURE_EXPORT_FILTERS,
            default_filter="PNG (*.png)"
        )

        if not file_path:
            return

        if sys.platform.startswith("win") and file_path.lower().startswith("c:\\program files"):
            QMessageBox.warning(
                self,
                "Permission Denied",
                "Windows does not allow saving files inside Program Files.\n"
                "Please choose another folder such as Documents or Desktop."
            )
            return

        try:
            # ✅ If user didn't type an extension, add the one from ext
            root, current_ext = os.path.splitext(file_path)
            if current_ext == "":
                ext = ext.lower().lstrip(".")
                file_path = f"{file_path}.{ext}"
            else:
                ext = current_ext.lower().lstrip(".")

            save_figure(fit_graph_figure(**self._graph), file_path, tight=True)
            QMessageBox.information(self, "Export Complete", f"Plot saved:\n{file_path}")
            self.status.showMessage("Export successful!", 3000)

        except Exception as e:
            QMessageBox.critical(self, "Export Failed", f"Could not save file:\n{e}")
            self.status.showMessage("Export failed!", 3000)

    def save_data(self):

        # All known e-Callisto station names
        station_list = [
            'ALASKA-ANCHORAGE', 'ALASKA-COHOE', 'ALASKA-HAARP', 'ALGERIA-CRAAG', 'ALMATY',
            'Arecibo-observatory', 'AUSTRIA-Krumbach', 'AUSTRIA-MICHELBACH', 'AUSTRIA-OE3FLB',
            'AUSTRIA-UNIGRAZ', 'Australia-ASSA', 'BRAZIL', 'BIR', 'Croatia-Visnjan', 'DENMARK',
            'EGYPT-Alexandria', 'EGYPT-SpaceAgency', 'ETHIOPIA', 'FINLAND-Siuntio', 'FINLAND-Kempele',
            'GERMANY-ESSEN', 'GERMANY-DLR', 'GLASGOW', 'GREENLAND', 'HUMAIN', 'HURBANOVO',
            'INDIA-GAURI', 'INDIA-Nashik', 'INDIA-OOTY', 'INDIA-UDAIPUR', 'INDONESIA',
            'ITALY-Strassolt', 'JAPAN-IBARAKI', 'KASI', 'KRIM', 'MEXART',
            'MEXICO-ENSENADA-UNAM', 'MEXICO-FCFM-UANL', 'MEXICO-FCFM-UNACH', 'MEXICO-LANCE-A',
            'MEXICO-LANCE-B', 'MEXICO-UANL-INFIERNILLO', 'MONGOLIA-UB', 'MRO', 'MRT1', 'MRT3',
            'Malaysia_Banting', 'NASA-GSFC', 'NORWAY-EGERSUND', 'NORWAY-NY-AALESUND', 'NORWAY-RANDABERG',
            'PARAGUAY', 'POLAND-BALDY', 'POLAND-Grotniki', 'ROMANIA', 'ROSWELL-NM', 'RWANDA',
            'SOUTHAFRICA-SANSA', 'SPAIN-ALCALA', 'SPAIN-PERALEJOS', 'SPAIN-SIGUENZA', 'SRI-Lanka',
            'SSRT', 'SWISS-CalU', 'SWISS-FM', 'SWISS-HB9SCT', 'SWISS-HEITERSWIL', 'SWISS-IRSOL',
            'SWISS-Landschlacht', 'SWISS-MUHEN', 'TAIWAN-NCU', 'THAILAND-Pathumthani', 'TRIEST',
            'TURKEY', 'UNAM', 'URUGUAY', 'USA-ARIZONA-ERAU', 'USA-BOSTON', 'UZBEKISTAN'
        ]

        # ✅ Extract Station
        station = "UNKNOWN"
        filename_lower = self.filename.lower()
        for s in station_list:
            if filename_lower.startswith(s.lower()):
                station = s
                break

        # ✅ Extract Date
        date_match = re.search(r'_(\d{4})(\d{2})(\d{2})_', self.filename)
        if date_match:
            date = f"{date_match.group(1)}-{date_match.group(2)}-{date_match.group(3)}"
        else:
            date = "UNKNOWN"

        # ✅ Excel File Handling
        from openpyxl import Workbook, load_workbook

        if self.existing_excel_checkbox.isChecked():
            path, _ = QFileDialog.getOpenFileName(self, "Select Existing Excel File", "", "Excel Files (*.xlsx)")
            if not path:
                return
            try:
                wb = load_workbook(path)
                ws = wb.active
            except Exception as e:
                QMessageBox.critical(self, "Load Error", f"Could not open Excel file:\n{str(e)}")
                return
        else:
            path, _ = QFileDialog.getSaveFileName(self, "Save as Excel", f"{self.filename}_data.xlsx",
                                                  "Excel Files (*.xlsx)")
            if not path:
                return
            try:
                wb = Workbook()
                ws = wb.active
                headers = [
                    "Date", "Station", "Best_fit", "R_sq", "RMSE",
                    "avg_freq", "avg_freq_err", "Avg_drift", "avg_drift_err",
                    "start_freq", "start_freq_err", "initial_shock_speed", "initial_shock_speed_err",
                    "initial_shock_height", "initial_shock_height_err", "avg_shock_speed", "avg_shock_speed_err",
                    "avg_shock_height", "avg_shock_height_err", "avg_drift_abs",
                    "density_model", "t0_s",
                ]
                ws.append(headers)
            except Exception as e:
                QMessageBox.critical(self, "Save Error", f"Could not create Excel file:\n{str(e)}")
                return

        # ✅ Extract and clean text
        def extract_val_err(label):
            # Remove HTML tags
            clean_text = re.sub(r'<[^>]+>', '', label.text())
            # Remove units and stray characters
            clean_text = re.sub(r'(MHz|km/s|Rₛ|s|/)', '', clean_text)
            # Clean spaces
            clean_text = clean_text.strip()
            # Extract value ± error
            value_text = clean_text.split(":")[-1].strip()
            if "±" in value_text:
                value, err = value_text.split("±")
                return value.strip(), err.strip()
            else:
                return value_text.strip(), ""

        # ✅ Read values
        try:
            best_fit = re.sub(r'<[^>]+>', '', self.equation_display.text()).replace("<sup>", "^").replace("</sup>", "")
            r2 = self.r2_display.text().split("=")[-1].strip()
            rmse = self.rmse_display.text().split("=")[-1].strip()

            avg_freq, avg_freq_err = extract_val_err(self.avg_freq_display)
            avg_drift, avg_drift_err = extract_val_err(self.drift_display)

            try:
                avg_drift_abs = abs(float(avg_drift))
            except ValueError:
                avg_drift_abs = ""

            start_freq, start_freq_err = extract_val_err(self.start_freq_display)
            init_speed, init_speed_err = extract_val_err(self.initial_shock_speed_display)
            init_height, init_height_err = extract_val_err(self.initial_shock_height_display)
            avg_speed, avg_speed_err = extract_val_err(self.avg_shock_speed_display)
            avg_height, avg_height_err = extract_val_err(self.avg_shock_height_display)

            row = [
                date, station, best_fit, r2, rmse,
                avg_freq, avg_freq_err, avg_drift, avg_drift_err,
                start_freq, start_freq_err, init_speed, init_speed_err,
                init_height, init_height_err, avg_speed, avg_speed_err,
                avg_height, avg_height_err, avg_drift_abs,
                density_model_label(self._selected_density_model()), float(self._fit_t0_s),
            ]

            ws.append(row)
            wb.save(path)
            self.status.showMessage("✅ Data saved to Excel successfully!", 3000)

        except Exception as e:
            QMessageBox.critical(self, "Write Error", f"Could not write to Excel file:\n{str(e)}")
            self.status.showMessage("❌ Failed to save data to Excel.", 3000)

    def plot_extra(self):
        choice = self.extra_plot_combo.currentText()
        freq_axis = np.asarray(getattr(self, "_shock_freq_values", getattr(self, "_fit_freq", self.freq)), dtype=float).reshape(-1)
        # choice -> (x, y, x label, y label, title suffix, on-screen colour)
        plots = {
            "Shock Speed vs Shock Height": (
                self.R_p, self.shock_speed, "Shock Height (Rₛ)", "Shock Speed (km/s)", "Shock_Speed_vs_Shock_Height", "green",
            ),
            "Shock Speed vs Frequency": (
                freq_axis, self.shock_speed, "Frequency (MHz)", "Shock Speed (km/s)", "Shock_Speed_vs_Frequency", "purple",
            ),
            "Shock Height vs Frequency": (
                self.R_p, freq_axis, "Shock Height (Rₛ)", "Frequency (MHz)", "Rs_vs_Freq", "red",
            ),
        }
        self.canvas.ax.clear()
        if choice in plots:
            x, y, x_label, y_label, suffix, colour = plots[choice]
            self.canvas.ax.scatter(x, y, color=colour, s=10)
            self.canvas.ax.set_xlabel(x_label)
            self.canvas.ax.set_ylabel(y_label)
            self.current_plot_title = f"{self.filename}_{suffix}"
            self.canvas.ax.set_title(self.current_plot_title)
            self._graph = {
                "x": x,
                "y": y,
                "data_label": "",
                "title": self.current_plot_title,
                "x_label": x_label,
                "y_label": y_label,
            }
            self.status.showMessage(f"{choice} plotted successfully!", 3000)
        self.canvas.ax.grid(True)
        self.canvas.draw()

    def closeEvent(self, event):
        self._emit_session_changed()
        super().closeEvent(event)
